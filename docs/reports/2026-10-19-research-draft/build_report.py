"""Build this dated research draft from public archived evidence only.

Use --figures-only with matplotlib/Pillow; then run without arguments with
python-docx/Pillow. Never loads private benchmark states or runs cryptography.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
SOURCE = HERE / '研究报告初稿.md'
DOCX = HERE / '研究报告初稿.docx'
COMPARISON = 'docs/research/2026-10-18-full-system-luo/evidence/comparison.json'
CARDS = 'docs/deliveries/2026-10-18-main-showcase-cards/evidence/cards.json'
PLACEMENT = 'docs/research/2026-10-19-full-system-placement/evidence/placement.json'
FOOLI = 'docs/research/2026-10-19-parameters/fooli-improvement-search.json'
BASE_COMMIT = 'b92967b92a94a4f8e18ae177ea299167b60aac67'
LABELS = ['MLDSA', 'Falcon', 'Luo']


def load(path):
    return json.loads((ROOT / path).read_text(encoding='utf-8'))


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def table(headers, rows):
    return '\n'.join(['| ' + ' | '.join(headers) + ' |',
                      '|' + '|'.join('---' for _ in headers) + '|'] +
                     ['| ' + ' | '.join(map(str, row)) + ' |' for row in rows])


def generate_evidence():
    comparison, cards, placement = load(COMPARISON), load(CARDS), load(PLACEMENT)
    assert len(comparison['profiles']) == len(cards['cards']) == 3
    assert not cards['selection']['final_three_representatives_selected']
    assert placement['source_commit'] == '3ae9826b10795bea1c11b8f848c7ac90d1047843'
    steps = [
        ('Setup', 'setup'), ('发送发钥', 'sender-key'), ('接收发钥', 'receiver-key'),
        ('首次加密', 'cold-encrypt'), ('首次解密含叶生成', 'first-decrypt-with-leaf-generation'),
        ('热加密', 'warm-encrypt'), ('热解密', 'warm-decrypt'),
        ('对象恢复', 'restore-objects'), ('恢复后解密', 'restored-decrypt'),
    ]
    timing = table(['操作', *LABELS], [
        [label, *(f"{p['timing_seconds'][key]['median']:.6f}" for p in comparison['profiles'])]
        for label, key in steps])
    size_keys = [
        ('完整密文', 'full_ciphertext'), ('公开状态', 'public_state'),
        ('主秘密状态', 'master_state'), ('发送状态', 'sender_state'),
        ('接收状态', 'receiver_state'),
    ]
    sizes = table(['对象', *LABELS], [[label, *(
        f"{c['measurements']['bytes'][key]['minimum']}–{c['measurements']['bytes'][key]['maximum']}"
        for c in cards['cards'])] for label, key in size_keys])
    variants = [('Windows仓库源码', 'source'), ('Windows同一wheel', 'windows_wheel'),
                ('WSL同一wheel', 'wsl_wheel')]
    placement_table = table(['运行位置', '完整加密', '完整解密', '每项观察次数'], [
        [name, f"{placement['timing_seconds'][key]['encrypt']['median'] * 1000:.3f}",
         f"{placement['timing_seconds'][key]['decrypt']['median'] * 1000:.3f}",
         placement['timing_seconds'][key]['encrypt']['count']]
        for name, key in variants])
    fooli = load(FOOLI)
    selected = []
    for n in (1024, 2048, 4096):
        candidates = [row for row in fooli['target_search']['all_cases']
                      if row['q'] == 2**127 - 1 and row['n'] == n and row['g'] == 16]
        selected.append(min(candidates, key=lambda row: row['m_bar']))
    fooli_table = table(['n', '完整维度d', '经典／量子（比特）', '本模型判断'], [
        [row['n'], f"{row['d']:,}",
         f"{row['core_svp_classical_bits']:.1f}／{row['core_svp_quantum_bits']:.1f}",
         '搜索下限' if row['core_svp_attack_blocksize'] == 53 else
         ('仅经典估计不低于128' if row['core_svp_classical_bits'] >= 128 else '经典估计不足128')]
        for row in selected])
    text = SOURCE.read_text(encoding='utf-8')
    for key, body in [('timing', timing), ('sizes', sizes), ('placement', placement_table), ('fooli', fooli_table)]:
        pattern = rf'(<!-- generated:{key}:start -->).*?(<!-- generated:{key}:end -->)'
        text, count = re.subn(pattern, lambda m: m[1] + '\n\n' + body + '\n\n' + m[2], text, flags=re.S)
        assert count == 1, key
    SOURCE.write_text(text, encoding='utf-8', newline='\n')
    fact = {
        'schema': 'IBME-RESEARCH-DRAFT/1', 'date': '2026-10-19', 'code_baseline': BASE_COMMIT,
        'new_cryptographic_runs': 0, 'private_state_files_opened': 0,
        'final_three_selected': False, 'security_claim': 'none',
        'backend_admitted': False, 'parameters_are_formal': False,
        'source_sha256': {p: sha(ROOT / p) for p in [COMPARISON, CARDS, PLACEMENT, FOOLI]},
        'fooli': {'selection': 'q=2**127-1, g=16; minimum m_bar in each declared n',
                  'source_schema_version': fooli['schema_version'], 'selected': selected,
                  'in_gate_search': fooli['in_gate_search']['stats'],
                  'core_svp_model': fooli['core_svp_model']},
        'profiles': [{
            'name': label, 'source_commit': p['source_commit'], 'receipt': p['file'],
            'security_group': p['security_group'], 'suite': p['suite'],
            'timing_seconds': {key: {k: p['timing_seconds'][key][k] for k in ('median', 'count')}
                               for _, key in steps},
            'bytes': c['measurements']['bytes'],
        } for label, p, c in zip(LABELS, comparison['profiles'], cards['cards'])],
        'placement': {
            'source_commit': placement['source_commit'], 'suite_digest': placement['suite_digest'],
            'suite': placement['suite'], 'independent_roots': placement['independent_roots'],
            'timing_seconds': {key: {op: {k: placement['timing_seconds'][key][op][k]
                                        for k in ('median', 'count')}
                                     for op in ('encrypt', 'decrypt')}
                               for _, key in variants},
        },
        'reference_pdf_sha256': {
            '有刺无刺+作品报告.pdf': 'dcd016ebda9cae48284a9bc42b239f8cf7dfe541d8839e680ee2d029710cc359',
            'crystal队作品设计报告.pdf': '38eb547782819462df449fb65e6b6376074000c49e7d91edc9ba76b67e9999a0',
        },
    }
    (HERE / 'report-facts.json').write_text(json.dumps(fact, ensure_ascii=False, indent=2) + '\n', encoding='utf-8', newline='\n')
    return comparison, placement


def figures(comparison, placement):
    import os
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.font_manager import FontProperties, fontManager
    from PIL import Image, ImageDraw, ImageFont

    directory = HERE / 'figures'
    directory.mkdir(exist_ok=True)
    font_path = Path('C:/Windows/Fonts/msyh.ttc') if os.name == 'nt' else Path('/mnt/c/Windows/Fonts/msyh.ttc')
    fontManager.addfont(str(font_path))
    prop = FontProperties(fname=str(font_path))
    plt.rcParams.update({'font.family': prop.get_name(), 'axes.unicode_minus': False,
                         'font.size': 10, 'axes.spines.top': False, 'axes.spines.right': False,
                         'figure.dpi': 160, 'savefig.dpi': 200})
    colors = ['#355c7d', '#687f6c', '#967343']
    fig, axes = plt.subplots(1, 2, figsize=(9.2, 3.5))
    for ax, key, title in zip(axes, ['warm-encrypt', 'warm-decrypt'], ['完整热加密', '完整热解密']):
        for i, p in enumerate(comparison['profiles']):
            data = p['timing_seconds'][key]
            ax.scatter([i + (j % 3 - 1) * .085 for j in range(9)], data['observations'],
                       s=25, color=colors[i], alpha=.8)
            ax.hlines(data['median'], i-.25, i+.25, color='black', lw=2)
        ax.set_xticks(range(3), LABELS)
        ax.set_title(title, fontproperties=prop, fontsize=12)
        ax.set_ylabel('时间 / 秒', fontproperties=prop)
        ax.grid(axis='y', alpha=.2); ax.set_axisbelow(True)
    fig.tight_layout(pad=1.5); fig.savefig(directory/'baseline.png'); plt.close(fig)
    fig, axes = plt.subplots(1, 2, figsize=(9.2, 3.7))
    keys = ['source', 'windows_wheel', 'wsl_wheel']
    for ax, op, title in zip(axes, ['encrypt', 'decrypt'], ['完整加密', '完整解密']):
        for i, key in enumerate(keys):
            data = placement['timing_seconds'][key][op]
            ax.boxplot([[x*1000 for x in data['observations']]], positions=[i], widths=.42,
                       patch_artist=True, boxprops={'facecolor': colors[i], 'alpha': .45},
                       medianprops={'color':'black'}, flierprops={'marker':'.'})
            ax.annotate(f"{data['median']*1000:.3f}", (i, data['median']*1000),
                        xytext=(0, 13), textcoords='offset points', ha='center', fontsize=9)
        ax.set_xticks(range(3), ['Windows\nsource', 'Windows\nwheel', 'WSL\nwheel'])
        ax.set_title(title, fontproperties=prop, fontsize=12)
        ax.set_ylabel('时间 / 毫秒', fontproperties=prop)
        ax.grid(axis='y', alpha=.2); ax.set_axisbelow(True)
    fig.tight_layout(pad=1.5); fig.savefig(directory/'placement.png'); plt.close(fig)

    font = lambda s: ImageFont.truetype(str(font_path), s)

    def canvas(w, h, title):
        img=Image.new('RGB',(w,h),'white'); draw=ImageDraw.Draw(img)
        draw.text((w/2,38),title,font=font(34),fill='#111111',anchor='mm')
        return img,draw

    def box(draw, xy, lines, fill='#f4f6f8'):
        draw.rounded_rectangle(xy, radius=12, fill=fill, outline='#89949f', width=2)
        x0,y0,x1,y1=xy
        for i,line in enumerate(lines):
            draw.text(((x0+x1)/2,(y0+y1)/2+(i-(len(lines)-1)/2)*40),line,
                      font=font(27),fill='#111111',anchor='mm')

    def arrow(draw, start, end):
        import math
        draw.line([start,end],fill='#46515a',width=3)
        a=math.atan2(end[1]-start[1],end[0]-start[0])
        pts=[end,(end[0]-17*math.cos(a-.5),end[1]-17*math.sin(a-.5)),
             (end[0]-17*math.cos(a+.5),end[1]-17*math.sin(a+.5))]
        draw.polygon(pts,fill='#46515a')

    img,d=canvas(1500,580,'双向身份匹配')
    box(d,(40,110,430,310),['发送者','实际身份 s','指定接收身份 r'])
    box(d,(570,110,970,310),['两层 HIBE 密文','父身份 r  →  子身份 s','IBS 认证接收目标与消息'])
    box(d,(1100,110,1460,310),['接收者','密钥身份 r′','期望发送身份 s′'])
    arrow(d,(430,210),(570,210));arrow(d,(970,210),(1100,210))
    box(d,(340,390,1160,520),['r = r′ 且 s = s′ 时恢复原消息','任意一侧错配 → 明确拒绝'],fill='#edf1ed')
    img.save(directory/'matching.png')
    img,d=canvas(1500,690,'加密封装与解密检查')
    xs=[40,420,800,1180]
    for x,lines in zip(xs,[['应用消息','与一次性公钥承诺'],['IBS 签名','接收身份与内部消息'],['Hybrid HIBE','种子胶囊与 GCM 载荷'],['BYTEAUTH','认证完整内层字节']]):
        box(d,(x,130,x+280,300),lines)
    for x in xs[:-1]:arrow(d,(x+280,215),(x+380,215))
    d.text((750,375),'解密按相反封装方向检查',font=font(29),fill='#111111',anchor='mm')
    box(d,(90,440,1410,620),['格式与外层一次性认证 → HIBE 与 GCM → IBS 验证',
                            '核对内部承诺 → 返回原消息或 BOTTOM'])
    img.save(directory/'layers.png')
    img,d=canvas(1500,760,'模块与配置数据流')
    box(d,(420,100,1080,210),['CLI 与已验收 Web 预设'])
    box(d,(350,290,1150,430),['公共装配与 Suite','版本参数绑定  容量预检  状态与规范编码'])
    arrow(d,(750,210),(750,290))
    box(d,(60,550,470,700),['HIBE 独立后端','参数与陷门表示封装'])
    box(d,(545,550,955,700),['IBS 独立后端','参数与签名编码封装'])
    box(d,(1030,550,1440,700),['公共协议与包装','双向匹配  Hybrid  BYTEAUTH'])
    for x in [265,750,1235]:arrow(d,(750,430),(x,550))
    img.save(directory/'architecture.png')
    # Readability crop only; all UI content is unchanged from archived evidence.
    with Image.open(ROOT/'docs/deliveries/2026-10-18-demo-acceptance/evidence/primary-success.jpg') as screenshot:
        assert screenshot.size == (622,2392)
        screenshot.crop((10,1030,612,1670)).save(directory/'demo-success-detail.png')


def build_docx():
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from urllib.parse import quote
    from PIL import Image

    doc = Document()
    sec=doc.sections[0]
    sec.page_width,sec.page_height=Cm(21),Cm(29.7)
    sec.top_margin=sec.bottom_margin=Cm(2.2)
    sec.left_margin=sec.right_margin=Cm(2.5)
    sec.header_distance=sec.footer_distance=Cm(1.1)
    sec.different_first_page_header_footer=True

    def setfont(style, cn, en, size, bold=False):
        style.font.name=en;style.font.size=Pt(size);style.font.bold=bold
        style.font.italic=False;style.font.underline=False
        style.font.color.rgb=RGBColor(0,0,0)
        rpr=style.element.get_or_add_rPr()
        fonts=rpr.find(qn('w:rFonts'))
        if fonts is None:fonts=OxmlElement('w:rFonts');rpr.insert(0,fonts)
        for key,value in [('eastAsia',cn),('ascii',en),('hAnsi',en)]:fonts.set(qn('w:'+key),value)
        for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme'):fonts.attrib.pop(qn('w:'+key),None)

    normal=doc.styles['Normal'];setfont(normal,'宋体','Times New Roman',11)
    normal.paragraph_format.line_spacing=1.35
    normal.paragraph_format.space_after=Pt(5)
    normal.paragraph_format.first_line_indent=Pt(22)
    normal.paragraph_format.widow_control=True
    for name,size in [('Title',23),('Subtitle',14),('Heading 1',17),('Heading 2',13),('Heading 3',11)]:
        style=doc.styles[name];setfont(style,'黑体','Times New Roman',size,name not in ('Subtitle',))
        style.paragraph_format.first_line_indent=Pt(0)
        style.paragraph_format.space_before=Pt(14)
        style.paragraph_format.space_after=Pt(10)
        style.paragraph_format.keep_with_next=True
    for name in ('Caption','Header','Footer'):
        setfont(doc.styles[name],'宋体','Times New Roman',9)
        doc.styles[name].paragraph_format.first_line_indent=Pt(0)
    setfont(doc.styles['List Bullet'],'宋体','Times New Roman',11)
    # The shipped Word template has a blue Title border; remove inherited rules.
    for style in doc.styles:
        for el in list(style.element.iter(qn('w:pBdr'))):
            el.getparent().remove(el)

    header=sec.header.paragraphs[0];header.style=doc.styles['Header']
    header.alignment=WD_ALIGN_PARAGRAPH.CENTER
    header.text='身份基匹配加密通用设计与优化实现  研究报告初稿'
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    foot.style=doc.styles['Footer']
    foot.add_run('第 ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    foot.add_run(' 页')
    doc.core_properties.title='基于多类格困难假设的抗量子身份基匹配加密通用设计与优化实现'
    doc.core_properties.subject='研究报告初稿 v0.1 依据2026年10月19日公开归档证据'
    doc.core_properties.author='IB-ME项目组'

    def inline(p, text):
        pos=0
        for match in re.finditer(r'\[([^\]]+)\]\(([^)]+)\)',text):
            p.add_run(text[pos:match.start()])
            path=(HERE/match[2]).resolve().relative_to(ROOT).as_posix()
            url=f'https://github.com/youzi258/IB-ME-from-HIBE-IBS/blob/{BASE_COMMIT}/{quote(path)}'
            link=OxmlElement('w:hyperlink')
            link.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
            run=OxmlElement('w:r');rpr=OxmlElement('w:rPr')
            color=OxmlElement('w:color');color.set(qn('w:val'),'000000');rpr.append(color)
            underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');rpr.append(underline)
            run.append(rpr);node=OxmlElement('w:t');node.text=match[1];run.append(node)
            link.append(run);p._p.append(link);pos=match.end()
        for n,part in enumerate(text[pos:].split('**')):
            r=p.add_run(part);r.bold=bool(n%2)
    lines=SOURCE.read_text(encoding='utf-8').splitlines()
    i=0
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line or line.startswith('<!--'):continue
        if line.startswith('# '):
            p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(95)
            inline(p,line[2:].replace('抗量子','抗量子\n',1));continue
        if line.startswith('## '):
            title=line[3:]
            p=doc.add_paragraph(title,style='Heading 1')
            if title in ('摘要','目录') or title.startswith(('第','附录','参考')):
                p.paragraph_format.page_break_before=True
            continue
        if line.startswith('### '):
            doc.add_paragraph(line[4:],style='Heading 2');continue
        if line.startswith('!['):
            m=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',line);assert m,line
            path=(HERE/m[2]).resolve()
            with Image.open(path) as im:w,h=im.size
            width=min(16.0, 11.5*w/h) if 'jpg' in path.suffix.lower() else 16.0
            if path.name=='demo-success-detail.png':width=13.2
            p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(path),width=Cm(width))
            caption=doc.add_paragraph(m[1],style='Caption');caption.alignment=WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.keep_with_next=False
            continue
        if line.startswith('```'):
            while i<len(lines) and not lines[i].startswith('```'):
                p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
                p.paragraph_format.left_indent=Cm(.4)
                p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.15
                r=p.add_run(lines[i]);r.font.name='Consolas';r.font.size=Pt(9)
                r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
                i+=1
            i+=1;continue
        if line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):
                rows.append(lines[i].strip());i+=1
            parsed=[[c.strip() for c in row.strip('|').split('|')] for row in rows]
            parsed=[r for r in parsed if not all(re.fullmatch('[: -]+',c) for c in r)]
            count=len(parsed[0]);t=doc.add_table(rows=0,cols=count)
            t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
            if count==4: widths=[3.6,4.3,4.1,4.0] if len(parsed)>10 else [5.4,3.5,3.5,3.6]
            else:widths=[4.0,5.7,6.3] if count==3 else [16/count]*count
            if parsed[0][0]=='操作' or parsed[0][0]=='对象':widths=[5.5,3.5,3.5,3.5]
            if parsed[0][0]=='编号':widths=[1.3,8.6,6.1]
            if parsed[0][0]=='路线':widths=[4.0,2.6,9.4]
            for c,w in zip(t.columns,widths):c.width=Cm(w)
            for rownum,row in enumerate(parsed):
                cells=t.add_row().cells
                for j,(cell,txt) in enumerate(zip(cells,row)):
                    cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    tcpr=cell._tc.get_or_add_tcPr()
                    margins=OxmlElement('w:tcMar')
                    for tag,val in [('top',60),('bottom',60),('left',95),('right',95)]:
                        el=OxmlElement('w:'+tag);el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa');margins.append(el)
                    tcpr.append(margins)
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9EDF0' if rownum==0 else 'FFFFFF');tcpr.append(shade)
                    p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0)
                    p.paragraph_format.line_spacing=1.18;p.paragraph_format.space_after=Pt(0)
                    p.paragraph_format.keep_with_next=rownum<len(parsed)-1
                    if rownum==0 or (len(txt)<25 and j>0):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    inline(p,txt)
                    for r in p.runs:r.font.size=Pt(9);r.bold=rownum==0
                trpr=t.rows[-1]._tr.get_or_add_trPr()
                no_split=OxmlElement('w:cantSplit');trpr.append(no_split)
                if rownum==0:
                    repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
            borders=OxmlElement('w:tblBorders')
            for edge in ('top','left','bottom','right','insideH','insideV'):
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            t._tbl.tblPr.append(borders)
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=.4
            continue
        if line.startswith(('表','算法')) and len(line)<65:
            p=doc.add_paragraph(line,style='Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True;continue
        p=doc.add_paragraph()
        if line.startswith('研究报告初稿') or re.fullmatch(r'\d{4}年\d+月\d+日',line):
            p.style='Subtitle';p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if line.startswith(('第1章','第2章','第3章','第4章','第5章','第6章','第7章','附录A','附录B','参考文献')):
            p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.space_after=Pt(10)
        if line.startswith('['):
            p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.line_spacing=1.15
            p.paragraph_format.space_after=Pt(5)
        inline(p,line)
        if line.startswith('['):
            for r in p.runs:r.font.size=Pt(10.5)
    doc.save(DOCX)
    print(f'Built {DOCX.name}: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables')


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--figures-only',action='store_true')
    args=parser.parse_args()
    comparison, placement=generate_evidence()
    if args.figures_only:figures(comparison, placement)
    else:build_docx()
